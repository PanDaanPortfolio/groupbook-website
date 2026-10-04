"""
Booyco Engineering — History File Migration Script
Extracts all records from the Excel History File and generates SQL INSERT statements
for the service_records table.

Ref: PD-BOOYCO-2026-001
Date: 25 March 2026
"""

import openpyxl
from datetime import datetime

wb = openpyxl.load_workbook('/mnt/project/Copy_of_History_File.xlsx')
ws = wb['Sheet1']

# ── LOOKUP MAPS ────────────────────────────────────────────────
# Reason for Action (Col R code -> name)
reason_map = {
    1: 'Poor Workmanship',
    2: 'User Induced Fault',
    3: 'Part Failure',
    4: 'Routine Maintenance',  # "Maintenance" in Excel = "Routine Maintenance" in new system
    5: 'Revision',
    6: 'Modification',
    7: 'Delivery',
    8: 'Concession',
}

# Repair or Warranty (Col U code -> name) — some rows have resolved text instead of codes
repair_war_map = {
    1: 'Repair (Charged)',
    2: 'Warranty Repair',
    3: 'Routine Maintenance',
    4: 'Checked Only',
    5: 'N/A',
    6: 'Commissioning',
}

# Text values that appear in resolved form
repair_war_text_map = {
    'Repair (Charged) (1)': 'Repair (Charged)',
    'Maintenance (3)': 'Routine Maintenance',
    'Warranty Repair(2)': 'Warranty Repair',
    'Warranty Repair (2)': 'Warranty Repair',
    'Checked': 'Checked Only',
    'Checked (4)': 'Checked Only',
    'N/A': 'N/A',
    'Commissioning': 'Commissioning',
}

# Failure Mode Main (Col AA code -> name)
failure_map = {
    0: 'None / N/A',
    1: 'Electrical',
    2: 'Mechanical Component',
    3: 'Structural',
    4: 'Gas Leak',
    5: 'None / N/A',  # Code 5 appears in data, treat as N/A
}

# Site code -> name
# Per Lauraine (25 Mar 2026):
#   Code 5 = Komatipoort (all Class 44 locos are Komatipoort)
#   Code 34 = "Frydensburg" was not a real site — tagged as Unassigned for Lauraine to reclassify
#   Code 39 = Leazonia (Drift Super Sand, DL16/DL18)
#   Code 47 = New Vaal (Sasolburg/Vereeniging area)
site_map = {
    1: 'Saldanha',
    2: 'Richardsbay',
    5: 'Komatipoort',
    7: 'Kimberley',
    10: 'Komatipoort',
    20: 'Unassigned',
    34: 'Unassigned',
    39: 'Leazonia',
    43: 'Unassigned',
    47: 'New Vaal',
}


def escape_sql(val):
    """Escape single quotes for SQL string literals."""
    if val is None:
        return "NULL"
    s = str(val).replace("'", "''").strip()
    if not s or s == "None":
        return "NULL"
    return f"'{s}'"


def resolve_action_type(reason_code, repair_war_val):
    """
    Determine the best action_type name from the two Excel columns.
    Priority: Repair/Warranty column if it gives a specific type,
    otherwise fall back to Reason for Action.
    """
    # Try repair/warranty first (more specific)
    if repair_war_val is not None:
        val_str = str(repair_war_val).strip()
        # Check if it's a text value
        if val_str in repair_war_text_map:
            return repair_war_text_map[val_str]
        # Check if it's a numeric code
        try:
            code = int(float(val_str))
            if code in repair_war_map:
                return repair_war_map[code]
        except (ValueError, TypeError):
            pass

    # Fall back to reason code
    if reason_code is not None:
        try:
            code = int(float(str(reason_code)))
            if code in reason_map:
                return reason_map[code]
        except (ValueError, TypeError):
            pass

    return None


def resolve_failure_mode(fm_code):
    """Resolve failure mode code to name."""
    if fm_code is None:
        return None
    try:
        code = int(float(str(fm_code)))
        return failure_map.get(code)
    except (ValueError, TypeError):
        # Might be a text value
        s = str(fm_code).strip()
        if 'Electrical' in s:
            return 'Electrical'
        if 'Component' in s:
            return 'Mechanical Component'
        if 'Structural' in s:
            return 'Structural'
        if 'Gas' in s:
            return 'Gas Leak'
        return None


# ── EXTRACT DATA ──────────────────────────────────────────────
records = []
skipped = 0

for r in range(15, ws.max_row + 1):
    # Check if row has any meaningful data
    serial = ws.cell(r, 1).value
    loco = ws.cell(r, 4).value
    date_val = ws.cell(r, 15).value
    site_code = ws.cell(r, 13).value
    comment = ws.cell(r, 31).value

    # Detect column shift: rows 807+ have SR in Col 22, and Col 24 = "Yes"/"No"/"N/A"
    col24_val = ws.cell(r, 24).value
    is_shifted = col24_val and str(col24_val).strip() in ('Yes', 'No', 'N/A')

    if is_shifted:
        # Shifted layout: Col 21=RepWar(text), Col 22=SR, Col 23=RepSuccess(code),
        # Col 24=RepSuccess(text), Col 25=FM(code), Col 26=FM(text)
        sr_no = ws.cell(r, 22).value
    else:
        # Normal layout: Col 24=SR
        sr_no = ws.cell(r, 24).value

    # Skip completely empty rows
    if not any([serial, loco, date_val, sr_no]):
        skipped += 1
        continue

    # Clean values (skip formula strings)
    def clean(val):
        if val is None:
            return None
        s = str(val).strip()
        if s.startswith("=") or s == "No Part No ENTERED!!" or s == "":
            return None
        return s

    # Parse date
    action_date = None
    if isinstance(date_val, datetime):
        action_date = date_val.strftime('%Y-%m-%d')
    elif date_val and not str(date_val).startswith("="):
        action_date = str(date_val)[:10]

    # Parse SR number
    sr_number = clean(sr_no)

    # Parse site — prioritise code-based lookup from site_map
    site_name = None
    if site_code is not None:
        try:
            code = int(float(str(site_code)))
            site_name = site_map.get(code)
        except (ValueError, TypeError):
            pass
    # Only fall back to resolved site value (Col 14) if code lookup failed
    if site_name is None:
        site_resolved = ws.cell(r, 14).value
        if site_resolved and not str(site_resolved).startswith("="):
            raw = str(site_resolved).strip()
            # Normalise: remove trailing code like "(2)" and match to known sites
            import re
            cleaned = re.sub(r'\s*\(\d+\)\s*$', '', raw).strip()
            # Try to match against site_map values
            matched = False
            for code, name in site_map.items():
                if cleaned.lower() == name.lower() or name.lower().startswith(cleaned.lower()):
                    site_name = name
                    matched = True
                    break
            if not matched:
                site_name = cleaned  # Use as-is if no match

    # Parse failure mode — handle column shift
    if is_shifted:
        fm_raw = ws.cell(r, 25).value       # FM code in shifted layout
        fm_resolved = ws.cell(r, 26).value   # FM text in shifted layout
    else:
        fm_raw = ws.cell(r, 27).value
        fm_resolved = ws.cell(r, 28).value
    failure_mode = resolve_failure_mode(fm_raw)
    if not failure_mode and fm_resolved and not str(fm_resolved).startswith("="):
        failure_mode = resolve_failure_mode(fm_resolved)

    # Parse action type — handle column shift
    reason_code = ws.cell(r, 18).value
    if is_shifted:
        repair_war = ws.cell(r, 21).value  # Already resolved text in shifted layout
    else:
        repair_war = ws.cell(r, 21).value
    action_type = resolve_action_type(reason_code, repair_war)

    # Parse repair success — handle column shift
    if is_shifted:
        repair_success_code = ws.cell(r, 23).value
    else:
        repair_success_code = ws.cell(r, 25).value

    # Parse other fields
    serial_number = clean(serial)
    unit_part = clean(ws.cell(r, 3).value)
    vehicle_ref = clean(loco)
    parts_no = clean(ws.cell(r, 8).value)
    parts_desc = clean(ws.cell(r, 9).value)
    parts_cost = ws.cell(r, 10).value
    comments = clean(comment)

    # Build parts_replaced text
    parts_text = None
    if parts_no or parts_desc:
        parts_parts = []
        if parts_no:
            parts_parts.append(parts_no)
        if parts_desc and not str(parts_desc).startswith("="):
            parts_parts.append(parts_desc)
        parts_text = " — ".join(parts_parts) if parts_parts else None

    # Repair success
    repair_success = None
    if repair_success_code is not None:
        try:
            code = int(float(str(repair_success_code)))
            repair_success = {1: 'Yes', 2: 'No', 3: 'N/A'}.get(code)
        except:
            s = str(repair_success_code).strip()
            if s in ('Yes', 'No', 'N/A'):
                repair_success = s

    records.append({
        'sr_number': sr_number,
        'action_date': action_date,
        'serial_number': serial_number,
        'unit_part_no': unit_part,
        'vehicle_ref': vehicle_ref,
        'site_name': site_name,
        'failure_mode': failure_mode,
        'action_type': action_type,
        'parts_replaced': parts_text,
        'comments': comments,
        'legacy_reason_code': reason_code if reason_code and not str(reason_code).startswith("=") else None,
        'legacy_repair_war': repair_war if repair_war and not str(repair_war).startswith("=") else None,
        'legacy_parts_cost': parts_cost if parts_cost and not str(parts_cost).startswith("=") else None,
        'legacy_repair_success': repair_success,
        'legacy_row': r,
    })


# ── GENERATE SQL ──────────────────────────────────────────────
print(f"-- Booyco History File Migration")
print(f"-- Records extracted: {len(records)}")
print(f"-- Rows skipped (empty): {skipped}")
print(f"-- Date range: {min(r['action_date'] for r in records if r['action_date'])} to {max(r['action_date'] for r in records if r['action_date'])}")
print(f"-- Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
print()
print("-- NOTE: site_id, failure_mode_id, and action_type_id are resolved via subqueries")
print("-- against the lookup tables seeded in booyco_schema.sql.")
print()

for i, rec in enumerate(records):
    site_sub = f"(SELECT id FROM sites WHERE name = {escape_sql(rec['site_name'])} LIMIT 1)" if rec['site_name'] else "NULL"
    fm_sub = f"(SELECT id FROM failure_modes WHERE name = {escape_sql(rec['failure_mode'])} LIMIT 1)" if rec['failure_mode'] else "NULL"
    at_sub = f"(SELECT id FROM action_types WHERE name = {escape_sql(rec['action_type'])} LIMIT 1)" if rec['action_type'] else "NULL"

    date_val = f"'{rec['action_date']}'" if rec['action_date'] else "NULL"

    legacy_reason = "NULL"
    if rec['legacy_reason_code'] is not None:
        try:
            legacy_reason = str(int(float(str(rec['legacy_reason_code']))))
        except:
            legacy_reason = "NULL"

    legacy_cost = "NULL"
    if rec['legacy_parts_cost'] is not None:
        try:
            legacy_cost = str(float(rec['legacy_parts_cost']))
        except:
            legacy_cost = "NULL"

    print(f"""INSERT INTO service_records (
    sr_number, action_date, serial_number, unit_part_no, vehicle_ref,
    site_id, failure_mode_id, action_type_id,
    parts_replaced, comments,
    legacy_reason_code, legacy_repair_war_code, legacy_parts_cost,
    legacy_repair_success, legacy_row_number
) VALUES (
    {escape_sql(rec['sr_number'])}, {date_val}, {escape_sql(rec['serial_number'])}, {escape_sql(rec['unit_part_no'])}, {escape_sql(rec['vehicle_ref'])},
    {site_sub}, {fm_sub}, {at_sub},
    {escape_sql(rec['parts_replaced'])}, {escape_sql(rec['comments'])},
    {legacy_reason}, {escape_sql(str(rec['legacy_repair_war']) if rec['legacy_repair_war'] else None)}, {legacy_cost},
    {escape_sql(rec['legacy_repair_success'])}, {rec['legacy_row']}
);""")

print()
print(f"-- Migration complete: {len(records)} records inserted.")
